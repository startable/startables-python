import os
import re
from pathlib import Path
from typing import Union

import pandas as pd

from startables import Bundle, ColumnMetadata, Table
from startables.startables import TableOrigin

TABLE_BLOCK_MARKER_PATTERN = r'^\*\*[^\*].*'


class WordTableBlockParsingError(Exception):
    pass


def import_from_word(path: Union[str, Path]) -> Bundle:
    """
    Imports table blocks from tables in an MS Word docx file.
    Only reads those Word tables whose top left cell is a table block start marker ('**table_name').

    MS Word documents cannot strictly live up to the StarTable format specification and therefore
    this function is not to be considered a formal reader. Rather, it is a utility that
    merely attempts to import table blocks.

    :param path: The path to the docx file
    :return: Bundle containing the imported tables
    """
    try:
        import docx
    except ImportError:
        raise ImportError("Missing optional dependency 'docx'. Install python-docx package "
                          "for MS Word support. Use pip or conda to install python-docx.") from None

    path = str(path)

    if not os.path.exists(path):
        raise IOError(f'File not found: {path}')

    word_doc = docx.Document(path)

    if not isinstance(word_doc, docx.document.Document):
        raise IOError(f'Not a docx Document: {word_doc}')

    tables = []
    for wt in word_doc.tables:

        # Does it even look like a StarTable table block?
        if not re.match(TABLE_BLOCK_MARKER_PATTERN, wt.cell(0, 0).text.strip()):
            continue
        if len(wt.rows) < 5:
            # Not enough rows for name, destinations, col names, col units, and at least one row of data.
            # Can't be a StarTable table block. Skip it.
            continue

        # Parse table
        table_name = wt.cell(0, 0).text.strip()[2:]
        destinations = set(wt.cell(1, 0).text.strip().split(' '))
        col_names = [cell.text.strip() for cell in wt.row_cells(2)]
        col_units = [cell.text.strip() for cell in wt.row_cells(3)]
        values = [[cell.text.strip() for cell in row.cells]
                  for row in wt.rows[4:]]

        try:
            df = pd.DataFrame(columns=col_names, data=values)
            col_specs = {n: ColumnMetadata(unit=u) for n, u in zip(col_names, col_units)}
            tables.append(Table(df, name=table_name, col_specs=col_specs,
                                destinations=destinations, origin=TableOrigin(path)))
        except AssertionError as e:
            # Malformed table
            raise WordTableBlockParsingError(
                f"Unable to parse table block '{table_name}' in document {path}") from e

    return Bundle(tables=tables, origin=TableOrigin(path))
